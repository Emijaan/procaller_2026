import re
from datetime import datetime

from django.db import transaction
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.access import agency_of, is_super_admin, require_perm, role_of, scoped_campaigns, scoped_leads, scoped_users
from accounts.models import Organization, User
from accounts.services import write_audit
from crm.models import Contact, ImportJob
from crm.serializers import ContactSerializer

from .models import Campaign
from .serializers import CampaignSerializer


def _scope(request):
    return scoped_campaigns(request.user, Campaign.objects.all())


def _reject_ibr(user, agency):
    if role_of(user) in {"user", "agent"}:
        return Response({"detail": "Users cannot run IBR"}, status=403)
    if not agency or not getattr(agency, "ibr_enabled", False):
        return Response({"detail": "IBR is disabled for this agency"}, status=403)
    return None


def _same_agency_manager(user):
    org = agency_of(user)
    if user.normalized_role == User.Role.MANAGER:
        return user
    boss = user.reports_to
    if boss and org and boss.organization_id == org.id and boss.normalized_role == User.Role.MANAGER:
        return boss
    return None


class CampaignListCreateView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        qs = _scope(request).select_related("agency", "manager", "admin", "created_by")
        return Response(CampaignSerializer(qs, many=True).data)

    def post(self, request):
        require_perm(request.user, "create_campaign")
        name = (request.data.get("name") or "").strip()
        if not name:
            return Response({"detail": "Campaign name is required"}, status=400)
        kind = (request.data.get("campaign_type") or Campaign.CampaignType.LIVE).strip()
        if kind not in {Campaign.CampaignType.LIVE, Campaign.CampaignType.IBR}:
            return Response({"detail": "Unknown campaign type"}, status=400)
        agency = agency_of(request.user)
        if is_super_admin(request.user) and request.data.get("agency"):
            agency = Organization.objects.filter(pk=request.data["agency"]).first()
        if kind == Campaign.CampaignType.IBR:
            denied = _reject_ibr(request.user, agency)
            if denied:
                return denied
        camp = Campaign.objects.create(
            name=name,
            description=request.data.get("description") or "",
            status=request.data.get("status") or Campaign.Status.DRAFT,
            campaign_type=kind,
            tts_prompt=request.data.get("tts_prompt") or "",
            tts_voice=request.data.get("tts_voice") or "",
            dial_method=request.data.get("dial_method") or Campaign.DialMethod.MANUAL,
            caller_id=request.data.get("caller_id") or "",
            caller_id_name=request.data.get("caller_id_name") or "ProCaller",
            code=request.data.get("code") or "",
            wrap_up_seconds=int(request.data.get("wrap_up_seconds") or 30),
            dial_ratio=max(1, min(32, int(request.data.get("dial_ratio") or 1))),
            agency=agency,
            created_by=request.user,
            admin=request.user if request.user.normalized_role == User.Role.ADMIN else None,
            manager=_same_agency_manager(request.user),
            active=True,
        )
        write_audit(request.user, "campaign_create", "campaign", camp.id, request)
        return Response(CampaignSerializer(camp).data, status=201)


class CampaignDetailView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        data = CampaignSerializer(camp).data
        data["lead_count"] = camp.contacts.count()
        return Response(data)

    def patch(self, request, pk):
        require_perm(request.user, "edit_campaign")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        next_kind = request.data.get("campaign_type", camp.campaign_type)
        if next_kind == Campaign.CampaignType.IBR or camp.campaign_type == Campaign.CampaignType.IBR:
            denied = _reject_ibr(request.user, camp.agency)
            if denied:
                return denied
        for field in (
            "name",
            "description",
            "status",
            "dial_method",
            "caller_id",
            "caller_id_name",
            "active",
            "code",
            "wrap_up_seconds",
            "dial_ratio",
            "auto_running",
            "campaign_type",
            "tts_prompt",
            "tts_voice",
        ):
            if field in request.data:
                value = request.data[field]
                if field == "wrap_up_seconds":
                    value = max(0, int(value or 30))
                if field == "dial_ratio":
                    value = max(1, min(32, int(value or 1)))
                if field == "auto_running":
                    value = value.lower() in {"1", "true", "yes"} if isinstance(value, str) else bool(value)
                setattr(camp, field, value)
        if "status" in request.data:
            camp.active = request.data["status"] in (Campaign.Status.ACTIVE, Campaign.Status.RUNNING)
        camp.save()
        if "assigned_users" in request.data:
            ids = []
            for raw in request.data.get("assigned_users") or []:
                try:
                    ids.append(int(raw))
                except (TypeError, ValueError):
                    continue
            camp.assigned_users.set(scoped_users(request.user).filter(pk__in=ids))
        write_audit(request.user, "campaign_update", "campaign", camp.id, request)
        return Response(CampaignSerializer(camp).data)


def _is_mp3(upload):
    name = (upload.name or "").lower()
    content = (getattr(upload, "content_type", "") or "").lower()
    return name.endswith(".mp3") or content in {"audio/mpeg", "audio/mp3", "audio/mpeg3"}


class CampaignHoldAudioView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        camp = _scope(request).filter(pk=pk).first()
        if not camp or not camp.hold_audio:
            return Response({"detail": "No hold audio uploaded for this campaign"}, status=404)
        from django.http import FileResponse

        return FileResponse(camp.hold_audio.open("rb"), as_attachment=False, filename="hold.mp3", content_type="audio/mpeg")

    def post(self, request, pk):
        require_perm(request.user, "edit_campaign")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Upload an MP3 file"}, status=400)
        if upload.size and upload.size > 5 * 1024 * 1024:
            return Response({"detail": "Hold audio must be 5MB or smaller"}, status=400)
        if not _is_mp3(upload):
            return Response({"detail": "Only MP3 files are allowed"}, status=400)
        if camp.hold_audio:
            camp.hold_audio.delete(save=False)
        camp.hold_audio = upload
        camp.save(update_fields=["hold_audio", "updated_at"])
        write_audit(request.user, "hold_audio_upload", "campaign", camp.id, request)
        return Response(CampaignSerializer(camp).data)

    def delete(self, request, pk):
        require_perm(request.user, "edit_campaign")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        if camp.hold_audio:
            camp.hold_audio.delete(save=True)
        write_audit(request.user, "hold_audio_delete", "campaign", camp.id, request)
        return Response(CampaignSerializer(camp).data)


def _digits(value):
    return re.sub(r"\D+", "", str(value or ""))


def _normalize_mobile(value):
    digits = _digits(value)
    if digits.startswith("91") and len(digits) == 12:
        return digits[2:]
    if digits.startswith("0") and len(digits) == 11:
        return digits[1:]
    return digits


def _read_excel(file_obj):
    from openpyxl import load_workbook

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("Excel file is empty")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    if not any(headers):
        raise ValueError("No column headers found")
    records = []
    for idx, row in enumerate(rows[1:], start=2):
        item = {}
        for i, header in enumerate(headers):
            if not header:
                continue
            val = row[i] if i < len(row) else None
            if isinstance(val, datetime):
                val = val.strftime("%Y-%m-%d")
            item[header] = "" if val is None else str(val).strip()
        if any(item.values()):
            item["_row"] = idx
            records.append(item)
    return headers, records


def _find_mobile_key(headers):
    for header in headers:
        key = header.lower().replace(" ", "")
        if key in {"mobile", "phone", "phonenumber", "mobileumber", "mobileno", "contact"}:
            return header
    return None


class CampaignImportPreviewView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        require_perm(request.user, "import_leads")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Upload an Excel file"}, status=400)
        name = (upload.name or "").lower()
        if not name.endswith((".xlsx", ".xls")):
            return Response({"detail": "Only .xlsx Excel files are allowed"}, status=400)
        if upload.size and upload.size > 15 * 1024 * 1024:
            return Response({"detail": "File is too large (max 15MB)"}, status=400)
        try:
            headers, records = _read_excel(upload)
        except Exception as exc:
            return Response({"detail": str(exc)}, status=400)
        mobile_key = _find_mobile_key(headers)
        if not mobile_key:
            return Response({"detail": "A Mobile column is required"}, status=400)
        existing = set(
            _normalize_mobile(p)
            for p in Contact.objects.filter(campaign=camp).values_list("phone", flat=True)
        )
        seen = set()
        valid, invalid, duplicates = [], [], []
        for rec in records:
            mobile = _normalize_mobile(rec.get(mobile_key))
            rec["_mobile"] = mobile
            if len(mobile) < 10:
                invalid.append({**rec, "_error": "Invalid mobile"})
                continue
            if mobile in seen or mobile in existing:
                duplicates.append({**rec, "_error": "Duplicate mobile"})
                continue
            seen.add(mobile)
            valid.append(rec)
        job = ImportJob.objects.create(
            agency=camp.agency or agency_of(request.user),
            campaign=camp,
            created_by=request.user,
            file_name=upload.name,
            status=ImportJob.Status.PREVIEW,
            mapping={"mobile": mobile_key, "columns": headers},
            total_rows=len(records),
            valid_rows=len(valid),
            invalid_rows=len(invalid),
            duplicate_rows=len(duplicates),
            errors=(invalid + duplicates)[:200],
        )
        write_audit(request.user, "excel_import_preview", "campaign", camp.id, request)
        return Response(
            {
                "job_id": job.id,
                "headers": headers,
                "mobile_column": mobile_key,
                "total_rows": len(records),
                "valid_rows": len(valid),
                "invalid_rows": len(invalid),
                "duplicate_rows": len(duplicates),
                "preview": valid[:25],
                "errors": (invalid + duplicates)[:50],
            }
        )


class CampaignImportConfirmView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        require_perm(request.user, "import_leads")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        job = ImportJob.objects.filter(pk=request.data.get("job_id"), campaign=camp).first()
        if not job:
            return Response({"detail": "Import preview not found. Upload the file again."}, status=400)
        upload = request.FILES.get("file")
        if not upload:
            return Response({"detail": "Upload the same Excel file to confirm import"}, status=400)
        headers, records = _read_excel(upload)
        mobile_key = job.mapping.get("mobile") or _find_mobile_key(headers)
        try:
            threads = max(1, min(32, int(request.data.get("dial_threads") or job.dial_threads or camp.dial_ratio or 1)))
        except (TypeError, ValueError):
            threads = max(1, min(32, int(camp.dial_ratio or 1)))
        skip_duplicates = request.data.get("skip_duplicates", True)
        existing = set(
            _normalize_mobile(p)
            for p in Contact.objects.filter(campaign=camp).values_list("phone", flat=True)
        )
        created = []
        errors = []
        seen = set()
        job.status = ImportJob.Status.IMPORTING
        job.save(update_fields=["status"])
        with transaction.atomic():
            for rec in records:
                mobile = _normalize_mobile(rec.get(mobile_key))
                if len(mobile) < 10:
                    errors.append({"row": rec.get("_row"), "error": "Invalid mobile", "mobile": rec.get(mobile_key)})
                    continue
                if mobile in seen or mobile in existing:
                    if skip_duplicates:
                        continue
                    errors.append({"row": rec.get("_row"), "error": "Duplicate mobile", "mobile": mobile})
                    continue
                extra = {k: v for k, v in rec.items() if k not in {mobile_key, "_row", "_mobile"}}
                name = extra.get("Name") or extra.get("name") or extra.get("Customer") or mobile
                company = extra.get("Company") or extra.get("company") or extra.get("Product") or ""
                email = ""
                for key, value in extra.items():
                    if key.lower().replace(" ", "") in {"email", "emailid", "mail"} and "@" in str(value):
                        email = str(value)[:254]
                        break
                contact = Contact(
                    name=str(name)[:160],
                    phone=mobile,
                    email=email,
                    company=str(company)[:160],
                    extra_data=extra,
                    campaign=camp,
                    agency=camp.agency or agency_of(request.user),
                    lead_status=Contact.LeadStatus.NEW,
                    owner=None,
                    assigned_admin=request.user if request.user.normalized_role == User.Role.ADMIN else camp.admin,
                )
                created.append(contact)
                seen.add(mobile)
                existing.add(mobile)
            Contact.objects.bulk_create(created, batch_size=500)
        job.status = ImportJob.Status.DONE
        job.imported_rows = len(created)
        job.dial_threads = threads
        job.errors = errors[:200]
        job.invalid_rows = len(errors)
        job.save()
        camp.dial_ratio = threads
        camp.save(update_fields=["dial_ratio", "updated_at"])
        write_audit(request.user, "excel_import", "campaign", camp.id, request, {"imported": len(created)})
        return Response(
            {
                "imported": len(created),
                "invalid": len(errors),
                "errors": errors[:50],
                "total_leads": camp.contacts.count(),
            }
        )


def _lead_board(camp, contacts):
    from telephony.models import Call
    from telephony.utils import format_duration

    from django.db.models import Count

    ids = [row.id for row in contacts]
    grouped = {pk: [] for pk in ids}
    totals = {}
    if ids:
        totals = {
            item["contact_id"]: item["c"]
            for item in Call.objects.filter(contact_id__in=ids).values("contact_id").annotate(c=Count("id"))
        }
        calls = (
            Call.objects.filter(contact_id__in=ids)
            .select_related("agent")
            .order_by("-id")
        )
        for call in calls:
            bucket = grouped.get(call.contact_id)
            if bucket is None or len(bucket) >= 8:
                continue
            bucket.append(
                {
                    "id": call.id,
                    "agent": call.agent.display_name if call.agent_id else "",
                    "state": call.state,
                    "outcome": call.outcome,
                    "disposition": call.disposition,
                    "duration": format_duration(call.duration_seconds),
                    "duration_seconds": call.duration_seconds,
                    "started_at": call.started_at,
                }
            )
    columns = []
    seen_cols = set()
    mobile_column = "Mobile"
    for job in camp.import_jobs.order_by("id"):
        mapping = job.mapping or {}
        if mapping.get("mobile"):
            mobile_column = mapping["mobile"]
            seen_cols.add(mapping["mobile"])
        for col in mapping.get("columns") or []:
            if col and col not in seen_cols:
                seen_cols.add(col)
                columns.append(col)
    for row in contacts:
        for key in row.extra_data or {}:
            if key and key not in seen_cols:
                seen_cols.add(key)
                columns.append(key)
    leads = []
    for row in contacts:
        history = grouped.get(row.id) or []
        latest = history[0] if history else None
        payload = ContactSerializer(row).data
        payload.update(
            {
                "attempt_count": totals.get(row.id, 0),
                "last_agent": (latest or {}).get("agent") or (row.owner.display_name if row.owner_id else ""),
                "last_outcome": (latest or {}).get("outcome") or "",
                "last_duration": (latest or {}).get("duration") or "00:00",
                "last_duration_seconds": (latest or {}).get("duration_seconds") or 0,
                "calls": history,
            }
        )
        leads.append(payload)
    return {
        "columns": columns,
        "mobile_column": mobile_column,
        "leads": leads,
    }


class CampaignLeadListView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        qs = scoped_leads(request.user, Contact.objects.filter(campaign=camp)).select_related("owner", "campaign")
        contacts = list(qs.order_by("id")[:500])
        return Response(_lead_board(camp, contacts))

    def post(self, request, pk):
        require_perm(request.user, "assign_leads")
        camp = _scope(request).filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        user_id = request.data.get("user_id")
        lead_ids = request.data.get("lead_ids") or []
        assignee = scoped_users(request.user).filter(pk=user_id).first()
        if not assignee:
            return Response({"detail": "User not found"}, status=404)
        qs = scoped_leads(request.user, Contact.objects.filter(campaign=camp, pk__in=lead_ids))
        updated = qs.update(owner=assignee, lead_status=Contact.LeadStatus.ASSIGNED)
        write_audit(request.user, "lead_assignment", "campaign", camp.id, request, {"count": updated})
        return Response({"assigned": updated})


class CampaignTtsView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        require_perm(request.user, "edit_campaign")
        camp = _scope(request).select_related("agency").filter(pk=pk).first()
        if not camp:
            return Response({"detail": "Not found"}, status=404)
        if camp.campaign_type != Campaign.CampaignType.IBR:
            return Response({"detail": "This campaign is not an IBR broadcast"}, status=400)
        denied = _reject_ibr(request.user, camp.agency)
        if denied:
            return denied
        if not (camp.tts_prompt or "").strip():
            return Response({"detail": "Write a speech prompt first"}, status=400)
        from telephony.tts import generate_campaign_audio

        try:
            result = generate_campaign_audio(camp)
        except ValueError as exc:
            return Response({"detail": str(exc)}, status=400)
        write_audit(request.user, "campaign_tts", "campaign", camp.id, request, result)
        return Response(result)
